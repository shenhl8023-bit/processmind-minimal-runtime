from contextlib import asynccontextmanager
import json

import pytest
from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.database import Base, configure_sqlite_engine
from app.models.models import Document, DocumentOperationDetail, Project
from app.services import llm_service
from app.services import route_rules_document_details_runtime as details_runtime
from app.services import route_rules_aggregation as aggregation
from app.services.route_rules_runtime import extract_route_set_with_llm
from app.services.document_operation_non_pdf_parser import parse_json_operation_details, parse_office_operation_details
from app.services.file_parser import extract_text
from app.services.process_tree_builder import build_superset_process_tree
from app.services.route_rules_parsing import _detail_row_normalized_names, _normalize_operation_name


CHINESE_ROUTE = {
    "零件名称": "验收轴", "材料": "45钢", "工艺路线": [
        {"工序号": 10, "工序名称": "下料", "工步内容": ["棒料锯切"]},
        {"工序号": 20, "工序名称": "车削", "工步内容": ["车外圆", "倒角"]},
        {"工序号": 30, "工序名称": "检验", "工步内容": ["检验尺寸"]},
    ],
}

@pytest.mark.parametrize("source_type", ["json_route_extract", "docx_llm_extract", "xlsx_llm_extract"])
def test_structured_source_steps_survive_tree_building_in_order_without_keyword_loss(source_type):
    details = [DocumentOperationDetail(id=1, project_id=1, document_id=1, pdf_name="route",
        operation_seq=10, operation_name="车削", normalized_name="车削",
        operation_content="1. 车外圆至20mm\n2. 倒角C1\n3. 车外圆至20mm", source_type=source_type),
        DocumentOperationDetail(id=2, project_id=1, document_id=1, pdf_name="route",
        operation_seq=20, operation_name="检验", normalized_name="检验",
        operation_content="1. 检查尺寸", source_type=source_type)]
    tree = build_superset_process_tree(operations=[], detail_rows=details, sample_count=1)
    assert tree[0]["step_items"] == ["车外圆至20mm", "倒角C1", "车外圆至20mm"]
    assert tree[1]["step_items"] == ["检查尺寸"]


@asynccontextmanager
async def document_db(tmp_path, filename, file_type):
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'non-pdf.db'}")
    configure_sqlite_engine(engine)
    sessions = async_sessionmaker(engine, expire_on_commit=False)
    try:
        async with engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        async with sessions() as db:
            db.add(Project(id=1, name="non-PDF"))
            await db.flush()
            db.add(Document(id=1, project_id=1, filename=filename, original_name=filename, file_type=file_type))
            await db.commit()
        yield sessions
    finally:
        await engine.dispose()


async def extract_details(db, project_id, *, docs=None):
    return await details_runtime.extract_document_operation_details(
        db, project_id, docs=docs, detail_row_normalized_names=_detail_row_normalized_names,
        normalize_operation_name=_normalize_operation_name,
    )


@pytest.mark.asyncio
async def test_chinese_json_route_produces_all_operations_and_steps(tmp_path, monkeypatch):
    monkeypatch.setattr(details_runtime, "UPLOAD_DIR", tmp_path)
    (tmp_path / "route.json").write_text(json.dumps(CHINESE_ROUTE, ensure_ascii=False), encoding="utf-8")
    async with document_db(tmp_path, "route.json", "json") as sessions:
        async with sessions() as db:
            rows = await details_runtime.ensure_document_operation_details(
                db, 1, extract_document_operation_details_fn=extract_details,
            )
            assert [(row.operation_seq, row.operation_name, row.operation_content) for row in rows] == [
                (10, "下料", "1. 棒料锯切"),
                (20, "车削", "1. 车外圆\n2. 倒角"),
                (30, "检验", "1. 检验尺寸"),
            ]
            assert all(row.source_type == "json_route_extract" for row in rows)


@pytest.mark.asyncio
async def test_param_json_dictionary_keeps_step_order_equipment_and_duplicate_stages(tmp_path, monkeypatch):
    monkeypatch.setattr(details_runtime, "UPLOAD_DIR", tmp_path)
    (tmp_path / "route.json").write_text(
        '{"车削":{"车外圆":"车外圆至20mm&&CA6140&&NULL","倒角":"倒角C1&&NULL&&NULL"},"检验":{"检验尺寸":"检验尺寸&&卡尺&&NULL"},"检验":{"终检":"终检外观&&NULL&&NULL"}}',
        encoding="utf-8",
    )
    async with document_db(tmp_path, "route.json", "json") as sessions:
        async with sessions() as db:
            rows = await extract_details(db, 1)
            assert [(row.operation_seq, row.operation_name) for row in rows] == [(10, "车削"), (20, "检验"), (30, "检验")]
            assert rows[0].operation_content == "1. 车外圆至20mm\n2. 倒角C1"
            assert rows[0].equipment_models == "CA6140"
            assert rows[1].equipment_models == "卡尺"
            assert rows[2].operation_content == "1. 终检外观"


@pytest.mark.asyncio
@pytest.mark.parametrize("file_type", ["docx", "xlsx"])
async def test_office_document_uses_actual_text_and_grounded_model_steps(tmp_path, monkeypatch, file_type):
    monkeypatch.setattr(details_runtime, "UPLOAD_DIR", tmp_path)
    path = tmp_path / f"route.{file_type}"
    if file_type == "docx":
        from docx import Document as WordDocument
        document = WordDocument()
        document.add_paragraph("10 车削 CA6140")
        document.add_paragraph("车外圆至20mm")
        document.add_paragraph("倒角C1")
        document.save(path)
    else:
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.append([10, "车削", "车外圆至20mm", "倒角C1", "CA6140"])
        workbook.save(path)
        workbook.close()

    async def model_response(system_prompt, user_prompt, **kwargs):
        assert "车外圆至20mm" in user_prompt
        assert "倒角C1" in user_prompt
        return json.dumps({"operations": [{"operation_seq": 10, "operation_name": "车削",
            "steps": ["车外圆至20mm", "倒角C1"], "equipment_types": "", "equipment_models": "CA6140"}]}, ensure_ascii=False)

    monkeypatch.setattr(llm_service, "call_llm", model_response)
    async with document_db(tmp_path, path.name, file_type) as sessions:
        async with sessions() as db:
            rows = await extract_details(db, 1)
            assert len(rows) == 1
            assert rows[0].operation_name == "车削"
            assert rows[0].operation_content == "1. 车外圆至20mm\n2. 倒角C1"
            assert rows[0].equipment_models == "CA6140"
            assert rows[0].source_type == f"{file_type}_llm_extract"


@pytest.mark.asyncio
@pytest.mark.parametrize("file_type", ["json", "doc", "docx", "xls", "xlsx"])
async def test_supported_non_pdf_details_survive_missing_source_and_reload(tmp_path, monkeypatch, file_type):
    monkeypatch.setattr(details_runtime, "UPLOAD_DIR", tmp_path)
    async with document_db(tmp_path, f"missing.{file_type}", file_type) as sessions:
        async with sessions() as db:
            db.add(DocumentOperationDetail(project_id=1, document_id=1, pdf_name=f"missing.{file_type}",
                operation_seq=10, operation_name="车削", normalized_name="车削", operation_content="1. 车外圆"))
            await db.commit()
        async with sessions() as db:
            rows = await details_runtime.ensure_document_operation_details(db, 1, extract_document_operation_details_fn=extract_details)
            assert [(row.operation_name, row.operation_content) for row in rows] == [("车削", "1. 车外圆")]
        async with sessions() as db:
            assert len((await db.execute(select(DocumentOperationDetail))).scalars().all()) == 1


@pytest.mark.asyncio
@pytest.mark.parametrize("empty_candidates", [False, True])
async def test_json_details_restore_exact_route_before_empty_or_noisy_summary_check(tmp_path, monkeypatch, empty_candidates):
    monkeypatch.setattr(details_runtime, "UPLOAD_DIR", tmp_path)
    monkeypatch.setattr(aggregation, "UPLOAD_DIR", tmp_path)
    (tmp_path / "route.json").write_text(json.dumps(CHINESE_ROUTE, ensure_ascii=False), encoding="utf-8")

    async def empty_summary(docs, **kwargs):
        return {}, ["route.json"], {}

    async def ensure(db, project_id, **kwargs):
        return await details_runtime.ensure_document_operation_details(
            db, project_id, extract_document_operation_details_fn=extract_details, **kwargs,
        )

    async with document_db(tmp_path, "route.json", "json") as sessions:
        async with sessions() as db:
            operations = await extract_route_set_with_llm(
                db, 1, collect_candidate_summary=empty_summary if empty_candidates else aggregation._collect_candidate_summary,
                ensure_document_operation_details_fn=ensure,
                merge_document_detail_rows_into_candidate_summary=aggregation._merge_document_detail_rows_into_candidate_summary,
                build_route_set_ops_from_candidates=aggregation.build_route_set_ops_from_candidates,
                set_extraction_task_state=lambda *_args, **_kwargs: {},
            )
            assert operations is not None
            assert [item["name"] for item in operations] == ["下料", "车削", "检验"]
            assert [item["seq"] for item in operations] == [10, 20, 30]
            assert len((await db.execute(select(DocumentOperationDetail))).scalars().all()) == 3


@pytest.mark.parametrize("payload", [
    {}, [], {"材料": "45钢"}, {"工艺路线": []}, {"工艺路线": [{"工序号": 10, "工序名称": "车削", "工步内容": []}]},
    {"工艺路线": [{"工序号": True, "工序名称": "车削", "工步内容": ["车外圆"]}]},
    {"工艺路线": [{"工序号": 10, "工序名称": "车削", "工步内容": [{"未知": "车外圆"}]}]},
    {"车削": {"车外圆": {"description": "车外圆"}}}, {"车削": {}},
])
def test_json_unknown_or_incomplete_structures_are_explicit_errors(payload):
    with pytest.raises(ValueError):
        parse_json_operation_details(json.dumps(payload, ensure_ascii=False))


@pytest.mark.parametrize("response", [
    "", "not json", "[]", '{"operations":[]}', '{"operations":null}',
    '{"operations":[{"operation_seq":10,"operation_name":"车削","steps":[],"equipment_types":"","equipment_models":""}]}',
    '{"operations":[{"operation_seq":10,"operation_name":"车削","steps":["磨削至Ra0.8"],"equipment_types":"","equipment_models":""}]}',
    '{"operations":[{"operation_seq":10,"operation_name":"车削","steps":["车外圆"],"equipment_types":"","equipment_models":"M999"}]}',
    '{"operations":[{"operation_seq":10,"operation_name":"车削","steps":"车外圆","equipment_types":"","equipment_models":""}]}',
])
def test_model_empty_malformed_or_fabricated_evidence_is_rejected(response):
    with pytest.raises(ValueError):
        parse_office_operation_details(response, "10 车削 车外圆 CA6140", "docx")


@pytest.mark.asyncio
async def test_invalid_json_batch_writes_no_partial_details_and_corrected_retry_recovers(tmp_path, monkeypatch):
    monkeypatch.setattr(details_runtime, "UPLOAD_DIR", tmp_path)
    first_path = tmp_path / "first.json"
    second_path = tmp_path / "second.json"
    first_path.write_text(json.dumps(CHINESE_ROUTE, ensure_ascii=False), encoding="utf-8")
    second_path.write_text('{"metadata":"unsupported"}', encoding="utf-8")
    async with document_db(tmp_path, first_path.name, "json") as sessions:
        async with sessions() as db:
            db.add(Document(id=2, project_id=1, filename=second_path.name, original_name=second_path.name, file_type="json"))
            await db.commit()
            with pytest.raises(HTTPException) as caught:
                await details_runtime.ensure_document_operation_details(db, 1, extract_document_operation_details_fn=extract_details)
            assert caught.value.status_code == 422
            assert "second.json" in caught.value.detail
            assert "工序" in caught.value.detail
            assert (await db.execute(select(DocumentOperationDetail))).scalars().all() == []
            second_path.write_text('{"清洗":{"冲洗":"冲洗干净&&NULL&&NULL"}}', encoding="utf-8")
            rows = await details_runtime.ensure_document_operation_details(db, 1, extract_document_operation_details_fn=extract_details)
            assert [row.operation_name for row in rows] == ["下料", "车削", "检验", "清洗"]


@pytest.mark.asyncio
@pytest.mark.parametrize("file_type", ["doc", "xls"])
async def test_legacy_office_binary_has_actionable_conversion_error(tmp_path, monkeypatch, file_type):
    monkeypatch.setattr(details_runtime, "UPLOAD_DIR", tmp_path)
    path = tmp_path / f"legacy.{file_type}"
    path.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"legacy binary")
    async with document_db(tmp_path, path.name, file_type) as sessions:
        async with sessions() as db:
            with pytest.raises(HTTPException) as caught:
                await extract_details(db, 1)
            assert caught.value.status_code == 422
            assert "另存" in caught.value.detail
            assert (await db.execute(select(DocumentOperationDetail))).scalars().all() == []


@pytest.mark.parametrize("sequences", [(20, 10, 20), (10, 10), (20, 10)])
def test_json_route_rejects_duplicate_or_decreasing_operation_numbers(sequences):
    payload = {"工艺路线": [
        {"工序号": sequence, "工序名称": "车削", "工步内容": ["车外圆"]}
        for sequence in sequences
    ]}
    with pytest.raises(ValueError, match="递增"):
        parse_json_operation_details(json.dumps(payload, ensure_ascii=False))


def office_operation(sequence, name, steps, machine=""):
    return {"operation_seq": sequence, "operation_name": name, "steps": steps,
        "equipment_types": "", "equipment_models": machine}


@pytest.mark.parametrize("source,operations", [
    ("10 车削 车外圆 CA6140\n20 磨削 磨孔 M1432", [office_operation(10, "车削", ["磨孔"], "M1432"), office_operation(20, "磨削", ["车外圆"], "CA6140")]),
    ("10 | 车削 | 车外圆 | CA6140\n20 | 磨削 | 磨孔 | M1432", [office_operation(10, "车削", ["车外圆"], "M1432"), office_operation(20, "磨削", ["磨孔"], "M1432")]),
    ("10\n车削\n车外圆\nCA6140\n20\n磨削\n磨孔\nM1432", [office_operation(10, "车削", ["磨孔"]), office_operation(20, "磨削", ["车外圆"])]),
    ("车削\n车外圆\nCA6140\n磨削\n磨孔\nM1432", [office_operation(10, "车削", ["磨孔"]), office_operation(20, "磨削", ["车外圆"])]),
    ("10 检验 检查毛坯\n20 车削 车外圆\n30 检验 检查尺寸", [office_operation(10, "检验", ["检查尺寸"]), office_operation(20, "车削", ["车外圆"]), office_operation(30, "检验", ["检查毛坯"])]),
    ("10 车削 车外圆 倒角\n20 检验 检查尺寸", [office_operation(10, "车削", ["倒角", "车外圆"]), office_operation(20, "检验", ["检查尺寸"])]),
    ("10 车削 车外圆 CA6140\n20 磨削 磨孔 M1432", [office_operation(10, "车削", ["磨孔"], "M1432")]),
    ("20 车削 车外圆\n30 检验 检查尺寸", [office_operation(10, "车削", ["车外圆"]), office_operation(20, "检验", ["检查尺寸"])]),
    ("车削\n车外圆\n磨削\n磨孔", [office_operation(10, "磨削", ["磨孔"]), office_operation(20, "车削", ["车外圆"])]),
])
def test_office_rejects_steps_equipment_or_sequence_assigned_to_wrong_operation(source, operations):
    with pytest.raises(ValueError):
        parse_office_operation_details(json.dumps({"operations": operations}, ensure_ascii=False), source, "docx")


@pytest.mark.parametrize("source,operations", [
    ("10 车削 CA6140\n车外圆\n倒角\n20 检验\n检查尺寸", [office_operation(10, "车削", ["车外圆", "倒角"], "CA6140"), office_operation(20, "检验", ["检查尺寸"])]),
    ("10 | 车削 | 车外圆 | CA6140\n20 | 磨削 | 磨孔 | M1432", [office_operation(10, "车削", ["车外圆"], "CA6140"), office_operation(20, "磨削", ["磨孔"], "M1432")]),
    ("10\n车削\n车外圆\n20\n磨削\n磨孔", [office_operation(10, "车削", ["车外圆"]), office_operation(20, "磨削", ["磨孔"])]),
    ("检验\n检查毛坯\n车削\n车外圆\n检验\n检查尺寸", [office_operation(10, "检验", ["检查毛坯"]), office_operation(20, "车削", ["车外圆"]), office_operation(30, "检验", ["检查尺寸"])]),
    ("工序号 | 10 | 工序名称 | 检验\n检查毛坯\n工序号 | 20 | 工序名称 | 车削\n车外圆\n工序号 | 30 | 工序名称 | 检验\n检查尺寸", [office_operation(10, "检验", ["检查毛坯"]), office_operation(20, "车削", ["车外圆"]), office_operation(30, "检验", ["检查尺寸"])]),
])
def test_office_binds_valid_numbered_table_and_unnumbered_repeated_operation_evidence(source, operations):
    rows = parse_office_operation_details(json.dumps({"operations": operations}, ensure_ascii=False), source, "xlsx")
    assert [(row.operation_seq, row.operation_name) for row in rows] == [(item["operation_seq"], item["operation_name"]) for item in operations]


def test_docx_mixed_paragraph_table_order_is_preserved_for_model_source(tmp_path):
    from docx import Document as WordDocument
    path = tmp_path / "mixed.docx"
    document = WordDocument()
    cells = document.add_table(rows=1, cols=2).rows[0].cells
    cells[0].text, cells[1].text = "车削", "车外圆"
    document.add_paragraph("检验")
    document.add_paragraph("检查尺寸")
    cells = document.add_table(rows=1, cols=2).rows[0].cells
    cells[0].text, cells[1].text = "淬火", "加热保温"
    document.save(path)

    source = extract_text(str(path), "docx", max_chars=None)
    assert source == "车削 | 车外圆\n检验\n检查尺寸\n淬火 | 加热保温"


@pytest.mark.asyncio
async def test_real_docx_cross_operation_error_is_atomic_then_corrected_retry_preserves_mixed_order(tmp_path, monkeypatch):
    from docx import Document as WordDocument
    monkeypatch.setattr(details_runtime, "UPLOAD_DIR", tmp_path)
    path = tmp_path / "mixed.docx"
    document = WordDocument()
    cells = document.add_table(rows=1, cols=3).rows[0].cells
    cells[0].text, cells[1].text, cells[2].text = "车削", "车外圆", "CA6140"
    document.add_paragraph("检验")
    document.add_paragraph("检查尺寸")
    cells = document.add_table(rows=1, cols=3).rows[0].cells
    cells[0].text, cells[1].text, cells[2].text = "磨削", "磨孔", "M1432"
    document.save(path)
    valid_operations = [
        office_operation(10, "车削", ["车外圆"], "CA6140"),
        office_operation(20, "检验", ["检查尺寸"]),
        office_operation(30, "磨削", ["磨孔"], "M1432"),
    ]
    wrong_operations = [office_operation(10, "车削", ["磨孔"], "M1432"), *valid_operations[1:]]
    attempts = iter([wrong_operations, valid_operations])

    async def model_response(system_prompt, user_prompt, **kwargs):
        assert user_prompt.index("车削") < user_prompt.index("检验") < user_prompt.index("磨削")
        return json.dumps({"operations": next(attempts)}, ensure_ascii=False)

    monkeypatch.setattr(llm_service, "call_llm", model_response)
    async with document_db(tmp_path, path.name, "docx") as sessions:
        async with sessions() as db:
            with pytest.raises(HTTPException) as caught:
                await details_runtime.ensure_document_operation_details(db, 1, extract_document_operation_details_fn=extract_details)
            assert caught.value.status_code == 422
            assert "对应工序" in caught.value.detail
            assert (await db.execute(select(DocumentOperationDetail))).scalars().all() == []
            rows = await details_runtime.ensure_document_operation_details(db, 1, extract_document_operation_details_fn=extract_details)
            assert [(row.operation_seq, row.operation_name) for row in rows] == [(10, "车削"), (20, "检验"), (30, "磨削")]
            assert [row.operation_content for row in rows] == ["1. 车外圆", "1. 检查尺寸", "1. 磨孔"]
